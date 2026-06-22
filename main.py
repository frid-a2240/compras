from fastapi import FastAPI, UploadFile, File, Request
from fastapi.responses import HTMLResponse, JSONResponse, FileResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from itsdangerous import URLSafeTimedSerializer, BadSignature, SignatureExpired
import pandas as pd
import os
import re
import shutil
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from database import get_cursor


app = FastAPI(root_path="/compras")
app.mount("/static", StaticFiles(directory="static"), name="static")
templates = Jinja2Templates(directory="templates")

@app.middleware("http")
async def strip_compras_prefix(request: Request, call_next):
    if request.scope["path"].startswith("/compras"):
        new_path = request.scope["path"][len("/compras"):] or "/"
        request.scope["path"] = new_path
        request.scope["raw_path"] = new_path.encode()
    return await call_next(request)

UPLOAD_DIR = "uploads"
HISTORIAL_DIR = "historial"
os.makedirs(UPLOAD_DIR, exist_ok=True)
os.makedirs(HISTORIAL_DIR, exist_ok=True)

# ── Sesiones ─────────────────────────────────────────────────
SECRET_KEY = "isp-compras-secret-2026"   # cámbialo en producción
serializer = URLSafeTimedSerializer(SECRET_KEY)

def crear_sesion(user_id: str) -> str:
    return serializer.dumps({"user": user_id})

def verificar_sesion(token: str) -> dict | None:
    try:
        return serializer.loads(token, max_age=28800)  # 8 horas
    except (BadSignature, SignatureExpired):
        return None

# ── Usuarios ──────────────────────────────────────────────────
USERS = {
    "203773": {"pass": "compras2026", "name": "Yudith Sayuri"},
    "204869": {"pass": "isp2026",     "name": "Johanna Terezhina"},
}

# ── Normalización de estatus ──────────────────────────────────
ESTATUS_MAP = {
    s.lower().strip(): "SURTIDO"
    for s in ["SURTIDO", "surtido", "Surtido", "SURTIDO ", "Surtido        ", "SURTIDO        ",
              "surtido        ", "SURTIDO   ", "surtido ", "sURTIDO", "suRTIDO"]
}
ESTATUS_MAP.update({
    s.lower().strip(): "PARCIAL"
    for s in ["PARCIAL", "Parcial", "parcial", "PARCIAL ", "Parcial        ", "PARCIAL        "]
})
ESTATUS_MAP.update({
    s.lower().strip(): "POR SURTIR"
    for s in ["POR SURTIR", "Por Surtir", "Por Surtir     ", "POR SURTIR     ", "poR SURTIR"]
})
ESTATUS_MAP["cancelada"] = "CANCELADA"


def normalize_estatus(val):
    if pd.isna(val):
        return "SIN ESTATUS"
    return ESTATUS_MAP.get(str(val).lower().strip(), str(val).strip().upper())


# ── Helpers de Estado Pago ───────────────────────────────────
def _norm_celda(val):
    """Devuelve la celda como string limpio en mayúsculas, o '' si está vacía."""
    if val is None:
        return ""
    try:
        if pd.isna(val):
            return ""
    except (TypeError, ValueError):
        pass
    s = str(val).strip().upper()
    return "" if s in ("NAN", "NAT", "NONE") else s


def _evaluar_celda(s):
    """Clasifica una sola celda: 'PAGADO' (cubre PAGADO y PAGADO X%), 'VACIO' u 'OTRO'."""
    if not s:
        return "VACIO"
    if s == "PAGADO":
        return "PAGADO"
    # "PAGADO 50%", "PAGADO 30%", "PAGADO 50%, COORDINAR ..." → también PAGADO
    if re.match(r"^PAGADO\s+\d+\s*%", s):
        return "PAGADO"
    return "OTRO"


# ── Clasificación de Estado Pago ─────────────────────────────
#   Regla: considera AMBAS columnas (OBSERVACIONES y PAGOS).
#   Si alguna dice "PAGADO" (con o sin porcentaje) → PAGADO
#   Las dos vacías → NO PAGADO
#   Cualquier otra cosa (folios, texto libre, fechas) → OTRO
#   El valor original de cada columna se conserva intacto para visualización.
def clasificar_pago(observaciones, pagos=None):
    obs = _norm_celda(observaciones)
    pag = _norm_celda(pagos)
    e_obs = _evaluar_celda(obs)
    e_pag = _evaluar_celda(pag)
    # Si alguna dice PAGADO (completo o con %) → PAGADO
    if "PAGADO" in (e_obs, e_pag):
        return "PAGADO"
    # Si las dos están vacías → NO PAGADO
    if e_obs == "VACIO" and e_pag == "VACIO":
        return "NO PAGADO"
    # Hay algo en al menos una pero no es PAGADO → OTRO
    return "OTRO"


def extraer_pct_pagado(observaciones, pagos=None):
    """Si alguna celda dice 'PAGADO X%' devuelve 'X%', si no, ''."""
    for v in (observaciones, pagos):
        s = _norm_celda(v)
        m = re.match(r"^PAGADO\s+(\d+)\s*%", s)
        if m:
            return f"{m.group(1)}%"
    return ""


def extraer_anio(valor):
    """Devuelve el año (str) de un valor de fecha o str de fecha, o '' si no se puede."""
    if valor is None or valor == "":
        return ""
    try:
        if pd.isna(valor):
            return ""
    except (TypeError, ValueError):
        pass
    try:
        dt = pd.to_datetime(valor, errors="coerce", dayfirst=True)   # ← dayfirst=True
        if pd.isna(dt):
            dt = pd.to_datetime(valor, errors="coerce", dayfirst=False)  # fallback ISO
        if pd.isna(dt):
            return ""
        anio = dt.year
        if 2000 <= anio <= 2100:
            return str(anio)
        return ""
    except Exception:
        return ""


def fmt_fecha_str(valor):
    """Formatea una fecha (datetime, string o lo que sea) a dd/mm/aaaa, o '' si no se puede."""
    if valor is None or valor == "":
        return ""
    try:
        if pd.isna(valor):
            return ""
    except (TypeError, ValueError):
        pass
    try:
        dt = pd.to_datetime(valor, errors="coerce", dayfirst=True)   # ← dayfirst=True
        if pd.isna(dt):
            dt = pd.to_datetime(valor, errors="coerce", dayfirst=False)
        if pd.isna(dt):
            return str(valor)[:10]
        return dt.strftime("%d/%m/%Y")
    except Exception:
        return str(valor)[:10]


_MESES_ES = ["ENERO","FEBRERO","MARZO","ABRIL","MAYO","JUNIO",
             "JULIO","AGOSTO","SEPTIEMBRE","OCTUBRE","NOVIEMBRE","DICIEMBRE"]

def extraer_mes(valor):
    """Devuelve el nombre del mes en español a partir de una fecha, o '' si no se puede."""
    if valor is None or valor == "":
        return ""
    try:
        if pd.isna(valor):
            return ""
    except (TypeError, ValueError):
        pass
    try:
        dt = pd.to_datetime(valor, errors="coerce", dayfirst=True)   # ← dayfirst=True
        if pd.isna(dt):
            dt = pd.to_datetime(valor, errors="coerce", dayfirst=False)
        if pd.isna(dt):
            return ""
        return _MESES_ES[dt.month - 1]
    except Exception:
        return ""


def load_requisiciones(path):
    df = pd.read_excel(path, sheet_name=None)
    sheet_name = [k for k in df.keys() if "REQ" in k.upper()][0]
    df = df[sheet_name].copy()
    df.columns = [str(c).strip() for c in df.columns]
    df = df[df["Requisicion"].notna()].copy()
    def parse_oc(x):
        if pd.isna(x) or str(x).strip() in ["", "nan"]:
            return ""
        s = str(x).strip().split("/")[0].split(",")[0].strip()
        try:
            return str(int(float(s)))
        except:
            return s
    df["Orden de compra"] = df["Orden de compra"].apply(parse_oc)
    df["Estatus Req %"] = df["Estatus Req %"].apply(
        lambda x: f"{int(float(x)*100)}%" if pd.notna(x) and str(x) not in ["", "nan"] else ""
    )
    df["Status de OC %"] = df["Status de OC %"].apply(
        lambda x: f"{int(float(x)*100)}%" if pd.notna(x) and str(x).strip() not in ["", "nan", "          "] else ""
    )
    return df


def load_oc(path):
    all_sheets = pd.read_excel(path, sheet_name=None)
    frames = []
    for sheet_name, df in all_sheets.items():
        if df.empty or "Hoja" in sheet_name:
            continue
        df.columns = [str(c).strip() for c in df.columns]
        oc_col   = next((c for c in df.columns if "compra" in c.lower() or "O_" in c), None)
        fecha_col = next((c for c in df.columns if "F_" in c or "FECHA" in c.upper()), None)
        prov_col  = next((c for c in df.columns if "PROV" in c.upper()), None)
        est_col   = next((c for c in df.columns if "statu" in c.lower() or "ESTATUS" in c.upper()), None)
        comp_col  = next((c for c in df.columns if "ompra" in c.lower() and "orden" not in c.lower() and "O_" not in c), None)
        total_col = next((c for c in df.columns if "otal" in c or "TOTAL" in c.upper()), None)
        fact_col  = next((c for c in df.columns if "FACTURA" in c.upper()), None)
        obs_col   = next((c for c in df.columns if "OBS" in c.upper()), None)
        pagos_col = next((c for c in df.columns if c.upper().strip() == "PAGOS"), None)
        if oc_col is None:
            continue
        sub = pd.DataFrame()
        def parse_oc_num(x):
            if pd.isna(x) or str(x).strip() in ["", "nan"]:
                return ""
            s = str(x).strip().split("/")[0].split(",")[0].strip()
            try:
                return str(int(float(s)))
            except:
                return s

        def fmt_celda(x):
            """Conserva la celda como string limpio. Números enteros sin .0, fechas como dd/mm/aaaa."""
            if pd.isna(x):
                return ""
            # Fechas
            if isinstance(x, (pd.Timestamp, datetime)):
                try:
                    return x.strftime("%d/%m/%Y")
                except Exception:
                    return str(x)[:10]
            # Enteros que pandas convirtió a float
            if isinstance(x, float):
                if x.is_integer():
                    return str(int(x))
                return str(x)
            return str(x).strip()

        def fmt_fecha(x):
            """Devuelve fecha como dd/mm/aaaa, o '' si no se puede parsear."""
            if pd.isna(x) or str(x).strip() in ["", "nan", "NaT"]:
                return ""
            try:
                dt = pd.to_datetime(x, errors="coerce", dayfirst=True)   # ← dayfirst=True
                if pd.isna(dt):
                    return ""
                return dt.strftime("%d/%m/%Y")
            except Exception:
                return str(x)[:10]

        sub["OC#"]         = df[oc_col].apply(parse_oc_num)
        sub["Fecha OC"]    = df[fecha_col].apply(fmt_fecha) if fecha_col else ""
        sub["Proveedor OC"]= df[prov_col].apply(lambda x: str(x).strip() if pd.notna(x) else "") if prov_col else ""
        sub["Estatus OC"]  = df[est_col].apply(normalize_estatus) if est_col else "SIN ESTATUS"
        sub["Comprador OC"]= df[comp_col].apply(lambda x: str(x).strip() if pd.notna(x) else "") if comp_col else ""
        sub["Total OC"]    = df[total_col].apply(lambda x: float(x) if pd.notna(x) and str(x).strip() not in ["","nan"] else 0.0) if total_col else 0.0
        sub["Factura"]     = df[fact_col].apply(lambda x: str(x).strip() if pd.notna(x) else "") if fact_col else ""
        sub["Observaciones"]= df[obs_col].apply(fmt_celda) if obs_col else ""
        sub["Pagos"]       = df[pagos_col].apply(fmt_celda) if pagos_col else ""
        sub["Hoja origen"] = sheet_name.strip()
        frames.append(sub[sub["OC#"] != ""])
    if not frames:
        return pd.DataFrame()
    result = pd.concat(frames, ignore_index=True)
    result = result.drop_duplicates(subset=["OC#"], keep="last")
    return result


# load_pagos queda definida por compatibilidad — actualmente NO se usa.
# La clasificación de pago ahora viene de OBSERVACIONES (ver clasificar_pago).
def load_pagos(path):
    try:
        df = pd.read_excel(path, sheet_name="PAGOS", header=1)
    except Exception:
        return {}
    df.columns = [str(c).strip() for c in df.columns]
    orden_col = next((c for c in df.columns if "ORDEN" in c.upper()), None)
    pago_col  = next((c for c in df.columns if c.upper().strip() == "PAGO"), None)
    if orden_col is None:
        return {}
    def parse_oc(x):
        if pd.isna(x) or str(x).strip() in ["", "nan"]:
            return ""
        s = str(x).strip().split("/")[0].split(",")[0].strip()
        try:
            return str(int(float(s)))
        except:
            return s
    pagos = {}
    for _, r in df.iterrows():
        oc = parse_oc(r.get(orden_col))
        if not oc:
            continue
        pago_val = r.get(pago_col) if pago_col else None
        tiene_pago = pd.notna(pago_val) and str(pago_val).strip() not in ["", "nan", "NaT"]
        fecha = ""
        if tiene_pago:
            try:
                fecha = pd.to_datetime(pago_val, dayfirst=True).strftime("%d/%m/%Y")
            except:
                fecha = str(pago_val)[:10]
        pagos[oc] = {"estado": "PAGADO" if tiene_pago else "PROGRAMADO", "fecha": fecha}
    return pagos


def make_match(df_req, df_oc):
    oc_dict = df_oc.set_index("OC#").to_dict("index") if not df_oc.empty else {}
    rows = []
    # 1. Recorrer requisiciones y vincular su OC (si la tiene)
    oc_referenciadas = set()
    for _, r in df_req.iterrows():
        oc_num  = str(r.get("Orden de compra", "")).strip()
        oc_data = oc_dict.get(oc_num, {})
        if oc_num and oc_data:
            oc_referenciadas.add(oc_num)
        # Fecha Req: primero la de la req, luego fallback a Fecha OC, luego Fecha entrada
        fecha_req_raw = r.get("Fecha de requisicion", "")
        fecha_oc_raw  = r.get("Fecha de OC", "") or oc_data.get("Fecha OC", "")
        fecha_ent_raw = r.get("Fecha entrada", "")
        fecha_req_fmt = fmt_fecha_str(fecha_req_raw) or fmt_fecha_str(fecha_oc_raw) or fmt_fecha_str(fecha_ent_raw)
        # Año: probar todas las fuentes hasta encontrar uno válido
        anio = (extraer_anio(fecha_req_raw)
                or extraer_anio(fecha_oc_raw)
                or extraer_anio(fecha_ent_raw))
        row = {
            "Requisicion":  r.get("Requisicion", ""),
            "Mes":          str(r.get("Mes", "")).strip() or extraer_mes(fecha_req_raw) or extraer_mes(fecha_oc_raw),
            "Año":          anio,
            "Fecha Req":    fecha_req_fmt,
            "Usuario":      str(r.get("Usuario", "")).strip(),
            "Departamento": str(r.get("Departamento", "")).strip(),
            "Proyecto":     str(r.get("Equipo / Proyecto", "")).strip(),
            "Descripcion":  str(r.get("Descripcion", "")).strip(),
            "Estatus Req":  r.get("Estatus Req %", ""),
            "Comprador":    str(r.get("Comprador", "")).strip(),
            "OC#":          oc_num,
            "Proveedor":    str(r.get("Proveedor", "")).strip() or oc_data.get("Proveedor OC", ""),
            "Total":        oc_data.get("Total OC", r.get("Total OC", 0)),
            "Estatus OC":   oc_data.get("Estatus OC", "SIN OC" if not oc_num else "NO ENCONTRADA"),
            "Factura":      oc_data.get("Factura", ""),
            "Fecha entrada":fmt_fecha_str(r.get("Fecha entrada", "")),
            "Observaciones":oc_data.get("Observaciones", ""),
            "Pagos":        oc_data.get("Pagos", ""),
        }
        # Match req ↔ OC
        if oc_num and oc_data:
            row["Match"] = "✓ CON OC"
        elif not oc_num:
            row["Match"] = "SIN OC"
        else:
            row["Match"] = "⚠ NO ENCONTRADA"

        # Estado Pago — considera AMBAS columnas (Observaciones + Pagos)
        if not oc_num or not oc_data:
            row["Estado Pago"] = "SIN OC"
            row["Fecha Pago"]  = ""
        else:
            row["Estado Pago"] = clasificar_pago(row["Observaciones"], row["Pagos"])
            row["Fecha Pago"]  = extraer_pct_pagado(row["Observaciones"], row["Pagos"])
        rows.append(row)

    # 2. Agregar OC huérfanas: OC que existen en el archivo de OC pero ninguna
    #    requisición las referencia. Aparecen como filas con Match = "SIN REQ".
    if not df_oc.empty:
        for _, oc in df_oc.iterrows():
            oc_num = str(oc.get("OC#", "")).strip()
            if not oc_num or oc_num in oc_referenciadas:
                continue
            obs   = oc.get("Observaciones", "")
            pagos = oc.get("Pagos", "")
            fecha_oc_str = fmt_fecha_str(oc.get("Fecha OC", ""))
            rows.append({
                "Requisicion":  "(SIN REQ)",
                "Mes":          extraer_mes(oc.get("Fecha OC", "")),
                "Año":          extraer_anio(oc.get("Fecha OC", "")),
                "Fecha Req":    fecha_oc_str,
                "Usuario":      "",
                "Departamento": "",
                "Proyecto":     "",
                "Descripcion":  "",
                "Estatus Req":  "",
                "Comprador":    str(oc.get("Comprador OC", "")).strip(),
                "OC#":          oc_num,
                "Proveedor":    str(oc.get("Proveedor OC", "")).strip(),
                "Total":        oc.get("Total OC", 0),
                "Estatus OC":   oc.get("Estatus OC", "SIN ESTATUS"),
                "Factura":      str(oc.get("Factura", "")).strip(),
                "Fecha entrada":fecha_oc_str,
                "Observaciones":obs,
                "Pagos":        pagos,
                "Match":        "SIN REQ",
                "Estado Pago":  clasificar_pago(obs, pagos),
                "Fecha Pago":   extraer_pct_pagado(obs, pagos),
            })
    return pd.DataFrame(rows)


# ═══════════════════════════════════════════════════════════════
#  RUTAS DE AUTENTICACIÓN
# ═══════════════════════════════════════════════════════════════

@app.get("/login", response_class=HTMLResponse)
async def login_page(request: Request):
    token = request.cookies.get("session")
    if token and verificar_sesion(token):
        return RedirectResponse("/", status_code=302)
    return templates.TemplateResponse("login.html", {"request": request})


@app.post("/auth")
async def auth(request: Request):
    body    = await request.json()
    user_id = body.get("user", "").strip()
    password= body.get("pass", "")
    user    = USERS.get(user_id)
    if not user:
        return JSONResponse({"ok": False, "msg": "Número de empleado no encontrado."}, status_code=401)
    if user["pass"] != password:
        return JSONResponse({"ok": False, "msg": "Contraseña incorrecta."}, status_code=401)
    token = crear_sesion(user_id)
    response = JSONResponse({"ok": True, "name": user["name"]})
    response.set_cookie(
        key="session",
        value=token,
        httponly=True,
        samesite="lax",
        max_age=28800,
    )
    return response


@app.get("/logout")
async def logout():
    response = RedirectResponse("/login", status_code=302)
    response.delete_cookie("session")
    return response


# ═══════════════════════════════════════════════════════════════
#  RUTAS PROTEGIDAS
# ═══════════════════════════════════════════════════════════════

def get_user(request: Request):
    token = request.cookies.get("session")
    if not token:
        return None
    data = verificar_sesion(token)
    return data["user"] if data else None


@app.get("/", response_class=HTMLResponse)
async def index(request: Request):
    if not get_user(request):
        return RedirectResponse("/login", status_code=302)
    return templates.TemplateResponse("index.html", {"request": request})


@app.post("/upload")
async def upload_files(request: Request, req_file: UploadFile = File(...), oc_file: UploadFile = File(...)):
    if not get_user(request):
        return JSONResponse({"error": "No autorizado"}, status_code=401)
    req_path = os.path.join(UPLOAD_DIR, "requisiciones.xlsx")
    oc_path  = os.path.join(UPLOAD_DIR, "ordenes.xlsx")
    with open(req_path, "wb") as f:
        shutil.copyfileobj(req_file.file, f)
    with open(oc_path, "wb") as f:
        shutil.copyfileobj(oc_file.file, f)
    return {"ok": True}


def guardar_historial(df, tipo, archivo="", user_id=""):
    """Inserta una entrada en el historial (PostgreSQL)."""
    total  = len(df)
    con_oc = int((df["Match"] == "✓ CON OC").sum())
    sin_oc = int((df["Match"] == "SIN OC").sum())
    try:
        with get_cursor() as cur:
            cur.execute(
                """
                INSERT INTO historial (tipo, archivo, total_req, con_oc, sin_oc, user_id)
                VALUES (%s, %s, %s, %s, %s, %s)
                """,
                (tipo, archivo, total, con_oc, sin_oc, user_id),
            )
    except Exception as e:
        # No tiramos la petición si la BD falla; solo lo loguea
        print(f"[historial] Error al guardar: {e}")

@app.get("/data")
async def get_data(request: Request):
    if not get_user(request):
        return JSONResponse({"error": "No autorizado"}, status_code=401)
    req_path = os.path.join(UPLOAD_DIR, "requisiciones.xlsx")
    oc_path  = os.path.join(UPLOAD_DIR, "ordenes.xlsx")
    if not os.path.exists(req_path) or not os.path.exists(oc_path):
        return JSONResponse({"error": "Archivos no cargados"}, status_code=400)
    df_req = load_requisiciones(req_path)
    df_oc  = load_oc(oc_path)
    df     = make_match(df_req, df_oc)
    guardar_historial(df, "PROCESADO", user_id=get_user(request) or "")
    total        = len(df)
    con_oc       = len(df[df["Match"] == "✓ CON OC"])
    sin_oc       = len(df[df["Match"] == "SIN OC"])
    no_encontrada= len(df[df["Match"] == "⚠ NO ENCONTRADA"])
    sin_req      = len(df[df["Match"] == "SIN REQ"])
    surtidas     = len(df[df["Estatus OC"] == "SURTIDO"])
    parciales    = len(df[df["Estatus OC"] == "PARCIAL"])
    por_surtir   = len(df[df["Estatus OC"] == "POR SURTIR"])
    # Conteos de pago
    pagados      = len(df[df["Estado Pago"] == "PAGADO"])
    no_pagados   = len(df[df["Estado Pago"] == "NO PAGADO"])
    pagos_otros  = len(df[df["Estado Pago"] == "OTRO"])
    df["Total"]  = pd.to_numeric(df["Total"], errors="coerce").fillna(0)
    total_monto  = float(df["Total"].sum())
    records = df.fillna("").to_dict("records")
    for r in records:
        r["Total"]      = f"${float(r['Total']):,.2f}" if r["Total"] else ""
        r["Requisicion"]= str(r["Requisicion"]).split(".")[0]
    return {
        "stats": {
            "total": total, "con_oc": con_oc, "sin_oc": sin_oc,
            "no_encontrada": no_encontrada, "sin_req": sin_req,
            "surtidas": surtidas, "parciales": parciales, "por_surtir": por_surtir,
            "pagados": pagados,
            "no_pagados": no_pagados, "pagos_otros": pagos_otros,
            "total_monto": f"${total_monto:,.2f}",
        },
        "rows": records
    }


@app.get("/export")
async def export_excel(request: Request):
    if not get_user(request):
        return JSONResponse({"error": "No autorizado"}, status_code=401)
    req_path = os.path.join(UPLOAD_DIR, "requisiciones.xlsx")
    oc_path  = os.path.join(UPLOAD_DIR, "ordenes.xlsx")
    if not os.path.exists(req_path) or not os.path.exists(oc_path):
        return JSONResponse({"error": "Archivos no cargados"}, status_code=400)
    df_req = load_requisiciones(req_path)
    df_oc  = load_oc(oc_path)
    df     = make_match(df_req, df_oc)
    ts       = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"COMPRAS_CONSOLIDADO_{ts}.xlsx"
    out_path = os.path.join(HISTORIAL_DIR, filename)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CONSOLIDADO"
    headers = ["Requisición","Mes","Fecha Req","Usuario","Departamento","Proyecto",
               "Descripción","Est. Req","Comprador","OC#","Proveedor",
               "Total","Estatus OC","Factura","Estado Pago",
               "Fecha Entrada","Observaciones","Pagos","Match"]
    header_fill = PatternFill("solid", fgColor="1B2A4A")
    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
    thin   = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )
    status_colors = {
        # Match
        "✓ CON OC":"E8F5E9","SIN OC":"FFF3E0","⚠ NO ENCONTRADA":"FFEBEE",
        "SIN REQ":"E3F2FD",
        # Estatus OC
        "SURTIDO":"C8E6C9","PARCIAL":"FFF9C4","POR SURTIR":"FFCCBC",
        # Estado Pago (nuevos)
        "PAGADO":"C8E6C9",
        "NO PAGADO":"FFCDD2",
        "OTRO":"E1BEE7",
    }
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill; cell.font = header_font
        cell.alignment = center; cell.border = thin
    ws.row_dimensions[1].height = 30
    col_widths = [12,10,12,22,22,20,24,10,14,10,30,14,14,12,14,14,24,24,16]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    for row_idx, (_, row) in enumerate(df.iterrows(), 2):
        total_val = row.get("Total", "")
        try:
            total_fmt = f"${float(total_val):,.2f}" if total_val not in ("", None) else ""
        except (ValueError, TypeError):
            total_fmt = ""
        values = [
            str(row.get("Requisicion","")).split(".")[0], row.get("Mes",""),
            row.get("Fecha Req",""), row.get("Usuario",""), row.get("Departamento",""),
            row.get("Proyecto",""), row.get("Descripcion",""), row.get("Estatus Req",""),
            row.get("Comprador",""), row.get("OC#",""), row.get("Proveedor",""),
            total_fmt, row.get("Estatus OC",""), row.get("Factura",""),
            row.get("Estado Pago",""),
            row.get("Fecha entrada",""), row.get("Observaciones",""),
            row.get("Pagos",""), row.get("Match",""),
        ]
        match_color = status_colors.get(row.get("Match",""), "FFFFFF")
        oc_color    = status_colors.get(row.get("Estatus OC",""), "FFFFFF")
        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name="Calibri", size=9); cell.border = thin
            if col_idx == 19:
                cell.fill = PatternFill("solid", fgColor=match_color); cell.alignment = center
            elif col_idx == 13:
                cell.fill = PatternFill("solid", fgColor=oc_color); cell.alignment = center
            elif col_idx == 15:
                cell.fill = PatternFill("solid", fgColor=status_colors.get(row.get("Estado Pago",""),"FFFFFF"))
                cell.alignment = center
            elif col_idx in [1,9,10,12]:
                cell.alignment = center
            else:
                cell.alignment = left
        if row_idx % 2 == 0 and row.get("Match","") not in status_colors:
            for col_idx in range(1, len(headers)):
                ws.cell(row=row_idx, column=col_idx).fill = PatternFill("solid", fgColor="F8F9FA")
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    # ── Hoja RESUMEN ───────────────────────────────────────────
    ws2 = wb.create_sheet("RESUMEN")
    ws2["A1"] = "RESUMEN EJECUTIVO"
    ws2["A1"].font = Font(bold=True, size=14, color="1B2A4A", name="Calibri")
    ws2["A3"] = "Generado:"; ws2["B3"] = datetime.now().strftime("%d/%m/%Y %H:%M")
    ws2["A5"] = "Total filas";         ws2["B5"] = len(df)
    ws2["A6"] = "Con OC vinculada";    ws2["B6"] = len(df[df["Match"]=="✓ CON OC"])
    ws2["A7"] = "Sin OC";              ws2["B7"] = len(df[df["Match"]=="SIN OC"])
    ws2["A8"] = "OC no encontrada";    ws2["B8"] = len(df[df["Match"]=="⚠ NO ENCONTRADA"])
    ws2["A9"] = "OC sin requisición";  ws2["B9"] = len(df[df["Match"]=="SIN REQ"])
    ws2["A11"]= "OC Surtidas";         ws2["B11"]= len(df[df["Estatus OC"]=="SURTIDO"])
    ws2["A12"]= "OC Parciales";        ws2["B12"]= len(df[df["Estatus OC"]=="PARCIAL"])
    ws2["A13"]= "OC Por surtir";       ws2["B13"]= len(df[df["Estatus OC"]=="POR SURTIR"])
    ws2["A15"]= "Pagadas";             ws2["B15"]= len(df[df["Estado Pago"]=="PAGADO"])
    ws2["A16"]= "No pagadas";          ws2["B16"]= len(df[df["Estado Pago"]=="NO PAGADO"])
    ws2["A17"]= "Pago otro estado";    ws2["B17"]= len(df[df["Estado Pago"]=="OTRO"])
    ws2["A19"]= "Monto total OC"
    df["Total_num"] = pd.to_numeric(df["Total"], errors="coerce").fillna(0)
    ws2["B19"] = f"${df['Total_num'].sum():,.2f}"
    ws2["B19"].font = Font(bold=True, color="1B6B3A", name="Calibri")
    ws2.column_dimensions["A"].width = 25; ws2.column_dimensions["B"].width = 20
    for row in ws2["A3:B19"]:
        for cell in row:
            cell.font = cell.font.copy(name="Calibri", size=10)
    wb.save(out_path)
    guardar_historial(df, "EXPORTADO", filename, user_id=get_user(request) or "")
    return FileResponse(out_path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=filename)


@app.get("/historial")
async def get_historial(request: Request):
    if not get_user(request):
        return JSONResponse({"error": "No autorizado"}, status_code=401)
    try:
        with get_cursor() as cur:
            cur.execute(
                """
                SELECT tipo, archivo, fecha, total_req, con_oc, sin_oc
                FROM historial
                ORDER BY fecha DESC
                LIMIT 50
                """
            )
            rows = cur.fetchall()
        # Formatear fecha igual que antes para no romper el frontend
        for r in rows:
            r["fecha"] = r["fecha"].strftime("%d/%m/%Y %H:%M:%S")
        return rows
    except Exception as e:
        print(f"[historial] Error al leer: {e}")
        return []


@app.get("/historial/descargar/{archivo}")
async def descargar_historial(archivo: str, request: Request):
    if not get_user(request):
        return JSONResponse({"error": "No autorizado"}, status_code=401)
    archivo = os.path.basename(archivo)
    ruta = os.path.join(HISTORIAL_DIR, archivo)
    if not os.path.exists(ruta) or not archivo.endswith(".xlsx"):
        return JSONResponse({"error": "Archivo no encontrado"}, status_code=404)
    return FileResponse(ruta,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=archivo)


@app.get("/files-loaded")
async def files_loaded(request: Request):
    if not get_user(request):
        return JSONResponse({"error": "No autorizado"}, status_code=401)
    req = os.path.exists(os.path.join(UPLOAD_DIR, "requisiciones.xlsx"))
    oc  = os.path.exists(os.path.join(UPLOAD_DIR, "ordenes.xlsx"))
    return {"req": req, "oc": oc, "ready": req and oc}